"""动态规划的阶段转移、折现与Gurobi独立最优性校验。"""

import unittest
import json
import re
import tempfile
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from openpyxl import load_workbook

from power_ga.data_collection.project_case import CasePlant, ProjectCase, align_stage_case
from power_ga.dynamic_planning.solver import _dp_tables, _gurobi_path, _state_costs, _validate_cases
from power_ga.genetic_iteration.project_solver import evaluate_project
from power_ga.operation_simulation.project_engine import simulate_project
from power_ga.result_output.dynamic_report import write_dynamic_report
from project_main import execute


class DynamicPlanningTests(unittest.TestCase):
    def test_stage_cost_progress_reports_intermediate_and_final_counts(self):
        """长时状态评价须报告单调的年度/总体进度，数值求值次数不变。"""
        messages = []
        calls = []

        def fake_evaluate(case, bits):
            """只构造已知二元状态费用；不运行8760小时以隔离进度逻辑。"""
            calls.append((case, bits))
            return sum(bits)

        with patch("power_ga.dynamic_planning.solver.evaluate_project",
                   side_effect=fake_evaluate):
            tables = _state_costs(("year_a", "year_b"), 5, messages.append)
        self.assertEqual(64, len(calls))
        self.assertEqual(2, len(tables))
        self.assertEqual(5, tables[0][31])
        counts = []
        for message in messages:
            match = re.search(r"状态 (\d+)/32， 总体 (\d+)/64 \(([\d.]+)%\)", message)
            self.assertIsNotNone(match, message)
            counts.append((int(match.group(1)), int(match.group(2)),
                           float(match.group(3))))
        self.assertGreater(len(counts), 2)  # 原实现只在两个阶段末各报告一次。
        self.assertLess(counts[0][1], 32)
        self.assertEqual((32, 64, 100.0), counts[-1])
        self.assertEqual(sorted(item[1] for item in counts),
                         [item[1] for item in counts])
        self.assertEqual(1, sum(item[1] == 64 for item in counts))

    def test_year_only_project_is_exogenous_and_charged(self):
        """只在中间年度出现的增量应供电并付年费，不能悄悄丢弃或挤占终期位。"""
        old = CasePlant(1, "合成煤电", "coal", 50, .4, 300, 0)
        extra = CasePlant(1, "合成煤电", "coal", 25, .2, 300, 1000)
        optional = CasePlant(2, "合成燃气", "gas", 15, 0, 400, 2000)
        final = ProjectCase(2025, 2027, (old,), (optional,), (75,) * 8760,
                            (0,) * 8760, 0, 657, (), "synthetic", "synthetic")
        annual = ProjectCase(2025, 2026, (old,), (extra,), (75,) * 8760,
                             (0,) * 8760, 0, 657, (), "synthetic", "synthetic")
        stage = align_stage_case(final, annual)
        run = simulate_project(stage, (0,), reserve_ratio=0)
        score = evaluate_project(stage, (0,))
        self.assertEqual(75, stage.existing[0].capacity_mw)
        self.assertAlmostEqual(1 / 3, stage.existing[0].min_fraction)
        self.assertEqual(1000, stage.existing[0].annual_cost_yuan)
        self.assertEqual(1000, score.annual_fixed_cost_yuan)
        self.assertEqual(0, run.summary["unserved"])
        self.assertLess(run.summary["max_balance_residual_mw"], 1e-6)
        self.assertIn("外生投运", " ".join(stage.warnings))

    def test_annual_existing_capacity_changes_hourly_feasibility(self):
        """前一年仅有50MW存量时不能偷用终期100MW；逐时缺口仍要守恒。"""
        def make(year, capacity):
            plant = CasePlant(1, "合成燃气", "gas", capacity, 0, 400, 0)
            return ProjectCase(2025, year, (plant,), (), (75,) * 8760,
                               (0,) * 8760, 0, 657, (), "synthetic", "synthetic")
        previous, terminal = make(2026, 50), make(2027, 100)
        stage = align_stage_case(terminal, previous)
        prior_run = simulate_project(stage, (), reserve_ratio=0)
        terminal_run = simulate_project(terminal, (), reserve_ratio=0)
        self.assertEqual(50, stage.existing[0].capacity_mw)
        self.assertEqual(25 * 8760, prior_run.summary["unserved"])
        self.assertEqual(0, terminal_run.summary["unserved"])
        self.assertLess(prior_run.summary["max_balance_residual_mw"], 1e-6)

    def test_stage_report_records_comparable_ga_path(self):
        """同一费用表的GA路径和差距必须写入阶段汇总，且不可行时不能报假差距。"""
        stage = lambda year, state, new, cost: SimpleNamespace(
            year=year, state=state, newly_built=new, annual_cost_yuan=cost,
            discounted_cost_yuan=cost, max_balance_residual_mw=0.0)
        result = SimpleNamespace(stages=(stage(2026, 0, 0, 10),
                                         stage(2027, 1, 1, 10)),
                                 objective_yuan=20.0, gurobi_objective_yuan=20.0,
                                 discount_rate=0.0,
                                 evaluations=({0: SimpleNamespace(total_cost_yuan=10,
                                                                    feasible=True)},
                                              {0: SimpleNamespace(total_cost_yuan=0,
                                                                    feasible=False),
                                               1: SimpleNamespace(total_cost_yuan=11,
                                                                    feasible=True)}))
        ga = SimpleNamespace(best=SimpleNamespace(feasible=True, total_cost_yuan=21.0),
                             evaluations=9, stop_reason="达到最大代数",
                             history=(SimpleNamespace(generation=0, best_cost_yuan=21.0,
                                      best_feasible=True, feasible_count=2,
                                      unique_count=3),))
        case = SimpleNamespace(candidates=(SimpleNamespace(name="合成燃气"),))
        with tempfile.TemporaryDirectory() as folder:
            exogenous = ((2026, SimpleNamespace(
                station_id=9, name="合成煤电", technology="coal",
                capacity_mw=25, annual_cost_yuan=1000)),)
            write_dynamic_report(result, case, Path(folder), .95, "source",
                                 (ga, (0, 1)), exogenous)
            audit = json.loads((Path(folder) / "dynamic_audit.json").read_text(encoding="utf-8"))
            self.assertEqual(.05, round(audit["ga_comparison"]["relative_gap"], 8))
            self.assertEqual("balanced_by_stage_v2", audit["ga_comparison"]["ga_seed_policy"])
            self.assertEqual("source", audit["stage_input_mode"])
            self.assertEqual(1000, audit["exogenous_annual_cost_yuan"])
            self.assertTrue((Path(folder) / "tables" / "ga_dp_comparison.csv").is_file())
            self.assertTrue((Path(folder) / "tables" / "ga_dp_stage_comparison.csv").is_file())
            book = load_workbook(Path(folder) / "tables" / "动态规划阶段汇总.xlsx",
                                 read_only=True, data_only=True)
            try:
                self.assertEqual(1000, book["年度外生计划"].cell(2, 6).value)
                self.assertEqual(1, book["逐年GA与DP"].cell(3, 9).value)
            finally:
                book.close()
        with tempfile.TemporaryDirectory() as folder:
            # 零费用最优与不可行候选都不能触发除零或产生虚假百分比。
            zero_result = SimpleNamespace(stages=result.stages, objective_yuan=0.0,
                                          gurobi_objective_yuan=0.0, discount_rate=0.0,
                                          evaluations=result.evaluations)
            infeasible_ga = SimpleNamespace(best=SimpleNamespace(feasible=False,
                                                total_cost_yuan=0.0),
                                            evaluations=1, stop_reason="达到最大代数",
                                            history=ga.history)
            write_dynamic_report(zero_result, case, Path(folder), .95,
                                 "source", (infeasible_ga, (0, 0)))
            audit = json.loads((Path(folder) / "dynamic_audit.json").read_text(encoding="utf-8"))
            self.assertIsNone(audit["ga_comparison"]["relative_gap"])

    def test_source_stage_curves_keep_terminal_candidate_encoding(self):
        """年度负荷/风光必须来自该年，建设容量仍按终期同一状态位解释。"""
        wind_final = CasePlant(10, "风电", "wind", 20, 0, 0, 0,
                               (.2,) * 8760)
        wind_year = CasePlant(10, "风电", "wind", 10, 0, 0, 0,
                              (.6,) * 8760)
        hydro_final = CasePlant(11, "水电", "hydro", 30, 0, 0, 0,
                                hydro_monthly_factor=(1.0,) * 12)
        hydro_year = CasePlant(11, "水电", "hydro", 25, 0, 0, 0,
                               hydro_monthly_factor=(0.3,) * 12)
        gas_final = CasePlant(20, "燃气", "gas", 15, 0.2, 400, 1000)
        final = ProjectCase(2025, 2027, (hydro_final,), (wind_final, gas_final),
                            (100,) * 8760, (5,) * 8760, .15, 800,
                            (), "source", "cost")
        annual = ProjectCase(2025, 2026, (hydro_year,), (wind_year,),
                             (70,) * 8760, (3,) * 8760, .12, 600,
                             (), "source", "cost")
        stage = align_stage_case(final, annual)
        self.assertEqual((70, 3, .6, 20), (stage.domestic_load_mw[0],
                                            stage.export_mw[0],
                                            stage.candidates[0].profile[0],
                                            stage.candidates[0].capacity_mw))
        self.assertEqual((15, 1000), (stage.candidates[1].capacity_mw,
                                       stage.candidates[1].annual_cost_yuan))
        self.assertEqual((25, .3), (stage.existing[0].capacity_mw,
                                   stage.existing[0].hydro_monthly_factor[0]))
        self.assertEqual(2026, stage.target_year)
        self.assertIn("未列电站20", " ".join(stage.warnings))

    def test_missing_annual_renewable_curve_is_rejected(self):
        """缺少该年可再生曲线时不能偷偷复用终期天气。"""
        wind = CasePlant(10, "风电", "wind", 20, 0, 0, 0,
                         (.2,) * 8760)
        final = ProjectCase(2025, 2027, (), (wind,), (1,) * 8760,
                            (0,) * 8760, .1, 1, (), "source", "cost")
        annual = ProjectCase(2025, 2026, (), (), (1,) * 8760,
                             (0,) * 8760, .1, 1, (), "source", "cost")
        with self.assertRaisesRegex(ValueError, "年度曲线"):
            align_stage_case(final, annual)

    def test_two_stage_known_path(self):
        """只允许0→1的建设转移；折现后应选择首年项目A、次年加B。"""
        def item(cost, feasible=True):
            return SimpleNamespace(total_cost_yuan=cost, feasible=feasible)
        # 两个候选项目按bit 0/1编码；状态3表示两者都投运。
        tables = ({0: item(20), 1: item(8), 2: item(14), 3: item(19)},
                  {0: item(30), 1: item(17), 2: item(16), 3: item(4)})
        path, cost = _dp_tables(tables, 0.10)
        gurobi_path, gurobi_cost = _gurobi_path(tables, 0.10, 2)
        self.assertEqual((1, 3), path)
        self.assertEqual(path, gurobi_path)
        self.assertAlmostEqual(8 + 4 / 1.1, cost)
        self.assertAlmostEqual(cost, gurobi_cost, places=6)

    def test_three_stage_known_path(self):
        """三阶段须允许0→1→1，并按每年的费用贴现而非重复投建。"""
        def item(cost):
            return SimpleNamespace(total_cost_yuan=cost, feasible=True)
        tables = ({0: item(1), 1: item(10)},
                  {0: item(12), 1: item(3)},
                  {0: item(12), 1: item(3)})
        expected = 1 + 3 / 1.1 + 3 / 1.1 ** 2
        self.assertEqual(((0, 1, 1), expected), _dp_tables(tables, .1))
        path, objective = _gurobi_path(tables, .1, 1)
        self.assertEqual((0, 1, 1), path)
        self.assertAlmostEqual(expected, objective)

    def test_full_horizon_uses_each_annual_source_and_reports_exogenous(self):
        """应用入口要串起三年真实8760小时阶段，并单列中间年外生计划。"""
        old = CasePlant(1, "合成燃气存量", "gas", 50, 0, 400, 0)
        option = CasePlant(2, "合成燃气候选", "gas", 40, 0, 400, 1000)
        extra = CasePlant(3, "合成煤电外生", "coal", 10, 0, 300, 500)
        def annual(year, load, candidates=()):
            return ProjectCase(2025, year, (old,), candidates, (load,) * 8760,
                               (0,) * 8760, 0, load * 8.76, (),
                               "synthetic", "synthetic")
        cases = {2026: annual(2026, 40),
                 2027: annual(2027, 80, (extra,)),
                 2028: annual(2028, 80, (option,))}
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            for year in (2026, 2027):
                workbook = root / str(year) / f"{year}.xlsx"
                workbook.parent.mkdir(parents=True)
                workbook.write_bytes(b"synthetic-input-marker")
            output = root / "result"
            def minimal_final_report(_case, _result, _simulation, dest, **_kwargs):
                """只隔离现有工作簿写入，保留入口后续清单与阶段报表路径。"""
                dest.mkdir(parents=True)
                (dest / "run_manifest.json").write_text(
                    json.dumps({"source_workbook_sha256": {}}, ensure_ascii=False),
                    encoding="utf-8")
            with patch("project_main.load_project_case",
                       side_effect=lambda _source, year, _cost: cases[year]), \
                 patch("project_main.write_project_report",
                       side_effect=minimal_final_report):
                _, result, simulation, _ = execute(
                    root, 2028, root / "cost.csv", output, figures=False,
                    algorithm="dp", planning_horizon="full",
                    population=12, generations=12)
                # 原双年度选项仍应只读取终期前一年，不受全规划年入口影响。
                pair_output = root / "pair_result"
                _, pair, _, _ = execute(
                    root, 2028, root / "cost.csv", pair_output, figures=False,
                    algorithm="dp", planning_horizon="pair",
                    population=12, generations=12)
            audit = json.loads((output / "dynamic_audit.json").read_text(encoding="utf-8"))
            manifest = json.loads((output / "run_manifest.json").read_text(encoding="utf-8"))
            self.assertEqual([2026, 2027, 2028], audit["years"])
            self.assertEqual([0, 1, 1], audit["states"])
            self.assertEqual([2027, 2028], [stage.year for stage in pair.stages])
            self.assertEqual(1, audit["exogenous_projects_count"])
            self.assertEqual(500, audit["exogenous_annual_cost_yuan"])
            self.assertEqual(0, audit["objective_yuan"] - audit["gurobi_objective_yuan"])
            self.assertEqual({"2026", "2027", "2028"}, set(manifest["stage_input_sha256"]))
            self.assertEqual({"population": 12, "generations": 12, "seed": 1},
                             manifest["ga_settings"])
            self.assertEqual("balanced_by_stage_v2", manifest["ga_seed_policy"])
            self.assertLess(simulation.summary["max_balance_residual_mw"], 1e-6)
            book = load_workbook(output / "tables" / "动态规划阶段汇总.xlsx",
                                 read_only=True, data_only=True)
            try:
                self.assertEqual(2027, book["年度外生计划"].cell(2, 1).value)
                self.assertEqual(4, book["逐年GA与DP"].max_row)
            finally:
                book.close()

    def test_full_horizon_rejects_single_factor_scenario(self):
        """不能用同一个负荷系数冒充多个独立年度的负荷和天气输入。"""
        option = CasePlant(2, "合成燃气候选", "gas", 1, 0, 400, 1000)
        final = ProjectCase(2025, 2028, (), (option,), (1,) * 8760,
                            (0,) * 8760, 0, 8.76, (), "synthetic", "synthetic")
        with patch("project_main.load_project_case", return_value=final):
            with self.assertRaisesRegex(ValueError, "完整年度序列必须读取各年度源库"):
                execute(Path("synthetic"), 2028, Path("cost.csv"),
                        Path("not-written"), figures=False, algorithm="dp",
                        planning_horizon="full", stage_input_mode="scaled")

    def test_batch_forwards_ga_comparison_settings(self):
        """EXE批处理须把种群、代数和种子传给同口径GA，不能偷偷使用默认值。"""
        from gridmind_app.main import main as app_main
        argv = ["gridmind", "--batch-source", "synthetic", "--batch-costs", "cost.csv",
                "--batch-output", "result", "--batch-algorithm", "dp",
                "--batch-planning-horizon", "full", "--batch-population", "24",
                "--batch-generations", "24", "--batch-seed", "11"]
        with patch("sys.argv", argv), patch("project_main.execute") as run:
            app_main()
        options = run.call_args.kwargs
        self.assertEqual((24, 24, 11), (options["population"],
                                       options["generations"], options["seed"]))
        self.assertEqual("full", options["planning_horizon"])

    def test_infeasible_initial_state_cannot_be_used(self):
        """不可行阶段状态应同时被Bellman和Gurobi排除。"""
        def item(cost, feasible=True):
            return SimpleNamespace(total_cost_yuan=cost, feasible=feasible)
        tables = ({0: item(1, False), 1: item(6)},
                  {0: item(1, False), 1: item(5)})
        self.assertEqual((1, 1), _dp_tables(tables, 0.07)[0])
        self.assertEqual((1, 1), _gurobi_path(tables, 0.07, 1)[0])

    def test_candidate_identity_cannot_shift_between_years(self):
        """同一状态位若在下一年换了技术或容量，必须在运行前拒绝。"""
        first = SimpleNamespace(target_year=2026, candidates=(
            SimpleNamespace(station_id=1, technology="gas", capacity_mw=10),))
        second = SimpleNamespace(target_year=2027, candidates=(
            SimpleNamespace(station_id=1, technology="gas", capacity_mw=12),))
        with self.assertRaisesRegex(ValueError, "容量"):
            _validate_cases((first, second), 0.07, 8)


if __name__ == "__main__":
    unittest.main()
