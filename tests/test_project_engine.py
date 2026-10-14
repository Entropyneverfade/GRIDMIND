"""公开合成8760小时场景：核对新项目引擎的储能连续性和能量平衡。"""

import unittest
from datetime import date, timedelta
from itertools import product
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import Workbook, load_workbook

from power_ga.common.domain import GASettings, PLANNING_MONTH_HOURS
from power_ga.data_collection.project_case import (TECH_LABELS, CasePlant, ProjectCase,
                                                    _make_plant, _read_curve, read_costs)
from power_ga.genetic_iteration.project_solver import evaluate_project, optimize_project
from power_ga.operation_simulation import project_engine
from power_ga.operation_simulation.project_engine import _hydro_target, _monthly_index, simulate_project
from power_ga.result_output.project_report import write_project_report


class ProjectEngineTests(unittest.TestCase):
    """使用完全合成的输入，不依赖或复制任何项目数据库。"""

    def test_wind_and_solar_share_surplus_curtailment(self):
        """风光均可用时按可用MW分摊弃电，且 Q风+Q光=需弃总量。"""
        wind, solar = project_engine._allocate_curtailment(80.0, 120.0, 50.0)
        self.assertAlmostEqual(wind, 20.0)
        self.assertAlmostEqual(solar, 30.0)
        self.assertAlmostEqual(wind + solar, 50.0)
        self.assertEqual(project_engine._allocate_curtailment(0.0, 120.0, 50.0),
                         (0.0, 50.0))
        self.assertEqual(project_engine._allocate_curtailment(80.0, 0.0, 50.0),
                         (50.0, 0.0))

    def test_eight_technologies_and_8760_balance(self):
        """八类技术可同时进入引擎，逐小时功率与储能状态残差接近零。"""
        daylight = tuple(0.7 if h % 24 in range(9, 16) else 0.0 for h in range(8760))
        wind = tuple(0.8 if h % 24 < 12 else 0.1 for h in range(8760))
        plants = (
            CasePlant(1, "Coal", "coal", 100, 0.5, 350, 0),
            CasePlant(2, "Gas", "gas", 30, 0.1, 600, 0),
            CasePlant(3, "Hydro", "hydro", 20, 0, 30, 0),
            CasePlant(4, "Biomass", "biomass", 10, 0.4, 400, 0),
            CasePlant(5, "Wind", "wind", 80, 0, 0, 0, wind),
            CasePlant(6, "PV", "solar", 50, 0, 0, 0, daylight),
            # 电池/抽蓄的内部容量分别为P×时长，往返效率按平方根分解。
            CasePlant(7, "Battery", "battery", 20, 0, 0, 0, (), 2, 1, 0.81),
            CasePlant(8, "Pumped", "pumped", 15, 0, 0, 0, (), 3, 1, 0.81),
        )
        case = ProjectCase(2025, 2029, plants, (), (100.0,) * 8760,
                           (5.0,) * 8760, 0.03, 0, (), "", "")
        run = simulate_project(case, (), keep_hourly=True)
        self.assertEqual(len(run.hourly), 8760)
        self.assertEqual(len(run.storage_hourly), 2 * 8760)
        self.assertLess(run.summary["max_balance_residual_mw"], 1e-7)
        self.assertGreater(run.summary["storage_final_mwh"], -1e-7)
        self.assertTrue(any(h.storage_charge_mw > 0 for h in run.hourly))
        self.assertTrue(any(h.storage_discharge_mw > 0 for h in run.hourly))
        # ΔE−sqrt(ηRT)Pch+Pdis/sqrt(ηRT)=0，对每站每小时成立。
        self.assertLess(max(abs(row[-1]) for row in run.storage_hourly), 1e-7)
        self.assertTrue(all(h.storage_charge_mw * h.storage_discharge_mw < 1e-9
                            for h in run.hourly))

    def test_small_storage_dispatches_first_for_peak_shaving(self):
        """两站都已充电且削峰只需2MW时，5MW小站先放电，20MW大站待机。"""
        wind_profile = (1.0,) + (0.0,) * 8759
        plants = (CasePlant(1, "Coal", "coal", 100, 0, 300, 0),
                  # 故意把大站列在前面，证明规则不依赖输入行次序。
                  CasePlant(2, "Large", "battery", 20, 0, 0, 0,
                            (), 2, 1, 0.81),
                  CasePlant(3, "Small", "battery", 5, 0, 0, 0,
                            (), 2, 1, 0.81),
                  CasePlant(4, "Wind", "wind", 120, 0, 0, 0, wind_profile))
        load = (80.0,) + (102.0,) * 8759
        case = ProjectCase(2025, 2029, plants, (), load, (0.0,) * 8760,
                           0.0, 0, (), "", "")
        run = simulate_project(case, (), keep_hourly=True, reserve_ratio=0.0)
        charged = {row[1]: row[3] for row in run.storage_hourly if row[0] == 0}
        discharged = {row[1]: row[4] for row in run.storage_hourly if row[0] == 1}
        self.assertGreater(charged["Large"], 0.0)
        self.assertGreater(charged["Small"], 0.0)
        self.assertAlmostEqual(discharged["Small"], 2.0)
        self.assertAlmostEqual(discharged["Large"], 0.0)
        self.assertLess(run.summary["max_balance_residual_mw"], 1e-7)
        self.assertLess(max(abs(row[-1]) for row in run.storage_hourly), 1e-7)

    def test_project_ga_matches_complete_two_bit_enumeration(self):
        """2位共4方案可穷举：只有新增煤电才满足可信容量，电池不应凭额定MW代替容量信用。"""
        existing = (CasePlant(1, "ExistingCoal", "coal", 90, 0, 300, 0),)
        candidates = (CasePlant(2, "NewCoal", "coal", 20, 0, 300, 1_000_000),
                      CasePlant(3, "Battery", "battery", 10, 0, 0, 2_000_000,
                                (), 2, 1, 0.81))
        case = ProjectCase(2025, 2029, existing, candidates, (100.0,) * 8760,
                           (0.0,) * 8760, 0.05, 0, (), "", "")
        exhaustive = [evaluate_project(case, chromosome)
                      for chromosome in product((0, 1), repeat=2)]
        optimum = min(exhaustive, key=lambda item: item.sort_key())
        self.assertEqual(optimum.chromosome, (1, 0))
        self.assertFalse(next(item for item in exhaustive
                              if item.chromosome == (0, 1)).feasible)
        settings = GASettings(population_size=8, max_generations=8,
                              min_generations=2, stall_generations=3,
                              elite_count=1, crossover_rate=0.8,
                              mutation_rate=0.10, seed=7)
        result = optimize_project(case, settings)
        self.assertEqual(result.best.chromosome, optimum.chromosome)
        self.assertAlmostEqual(result.best.total_cost_yuan, optimum.total_cost_yuan)
        self.assertTrue(result.best.feasible)

    def test_dry_hydro_does_not_replace_daily_firm_commitment(self):
        """水量为零的50MW水电不计日前全天可用容量，燃气须开机补足负荷与备用。"""
        plants = (CasePlant(1, "DryHydro", "hydro", 50, 0, 0, 0,
                            hydro_monthly_factor=(0.0,) * 12),
                  CasePlant(2, "Coal", "coal", 70, 0, 300, 0),
                  CasePlant(3, "Gas", "gas", 60, 0, 500, 0))
        case = ProjectCase(2025, 2029, plants, (), (100.0,) * 8760,
                           (0.0,) * 8760, 0.03, 0, (), "", "")
        run = simulate_project(case, (), keep_hourly=True)
        self.assertAlmostEqual(run.summary["unserved"], 0.0)
        self.assertAlmostEqual(run.summary["reserve_shortfall"], 0.0)
        self.assertAlmostEqual(run.hourly[0].hydro_mw, 0.0)
        self.assertGreater(run.hourly[0].gas_mw, 0.0)

    def test_hydro_budget_is_upper_bound_with_zero_net_load(self):
        """整月无正净负荷时不为用完来水而发电，其余月份仍可按需用水。"""
        net = [0.0] * 8760
        net[31 * 24] = 10.0  # 2月首小时留有正净负荷，检验非零月份仍调度。
        budget = [5.0] * 12
        target = _hydro_target(net, 20.0, budget)
        self.assertEqual(len(target), 8760)
        self.assertAlmostEqual(sum(target[:31 * 24]), 0.0)
        self.assertAlmostEqual(target[31 * 24], 5.0, places=7)
        self.assertAlmostEqual(sum(target), 5.0, places=7)

    def test_vector_hydro_matches_scalar_equation(self):
        """数组加速后仍满足原式 p=min(Pmax,λw) 与每月水量上限。"""
        weights = [float((hour * 17) % 73) for hour in range(8760)]
        budgets = [5000.0] * 12
        output = _hydro_target(weights, 20.0, budgets)
        first = 0
        for size in PLANNING_MONTH_HOURS:
            part = weights[first:first + size]
            low, high = 0.0, max(1.0, 20.0 / min(v for v in part if v > 0))
            for _ in range(60):
                middle = (low + high) / 2
                if sum(min(20.0, middle * value) for value in part) < 5000.0:
                    low = middle
                else:
                    high = middle
            self.assertLess(max(abs(actual - min(20.0, high * value))
                                for actual, value in zip(output[first:first + size], part)), 1e-10)
            self.assertLessEqual(sum(output[first:first + size]), 5000.0 + 1e-7)
            first += size

    def test_leap_year_uses_8760_hour_planning_calendar(self):
        """2028年省略2月29日；水电运行、逐月工作簿和窗口共用365日边界。"""
        self.assertEqual(sum(PLANNING_MONTH_HOURS), 8760)
        self.assertEqual(_monthly_index()[744:1416], [1] * 672)
        plants = (CasePlant(1, "Hydro", "hydro", 2, 0, 30, 0),
                  CasePlant(2, "Coal", "coal", 10, 0, 300, 0))
        with TemporaryDirectory() as folder:
            root = Path(folder)
            # 报告只对源文件留SHA256；合成标记文件不会充作项目工作簿输入。
            for year in (2025, 2028):
                source = root / str(year)
                source.mkdir()
                (source / f"{year}.xlsx").write_bytes(b"synthetic-test-source")
            cost = root / "cost.csv"
            cost.write_text("synthetic-test-cost", encoding="utf-8")
            case = ProjectCase(2025, 2028, plants, (), (5.0,) * 8760,
                               (0.0,) * 8760, 0.0, 0, (), str(root), str(cost))
            run = simulate_project(case, (), keep_hourly=True)
            self.assertEqual(len(run.hourly), 8760)
            self.assertLess(run.summary["max_balance_residual_mw"], 1e-7)
            write_project_report(case, None, run, root / "result", make_figures=False)
            book = load_workbook(root / "result/tables/全年规划运行结果汇总.xlsx", read_only=True)
            rows = list(book["逐月统计"].values)
            book.close()
            # 2月起止为744..1415，3月从1416开始；Σ月负荷应等于全年负荷。
            self.assertEqual((rows[2][1], rows[2][2], rows[3][1]), (744, 1415, 1416))
            self.assertAlmostEqual(sum(row[6] for row in rows[1:]), 5.0 * 8760)

    def test_python_figures_need_no_origin_and_keep_hourly_values(self):
        """无 Origin 后端从已存 CSV 生成峰日/8760图，并能独立重绘。"""
        from power_ga.result_output.python_plot import render_archived_figure

        with TemporaryDirectory() as folder:
            root = Path(folder)
            for year in (2025, 2029):
                (root / str(year)).mkdir()
                (root / str(year) / f"{year}.xlsx").write_bytes(b"synthetic-only")
            (root / "cost.csv").write_text("synthetic-cost", encoding="utf-8")
            plants = (CasePlant(1, "Coal", "coal", 10, 0, 300, 0),)
            case = ProjectCase(2025, 2029, plants, (), (6.0,) * 8760,
                               (0.0,) * 8760, 0.0, 0, (), str(root), str(root / "cost.csv"))
            run = simulate_project(case, (), keep_hourly=True)
            report = root / "report"
            write_project_report(case, None, run, report, figure_backend="python")
            for chart in ("02_峰荷日电力平衡", "05_全年8760电力平衡"):
                figure = report / "figures" / chart
                self.assertTrue((figure / "figure.png").is_file())
                self.assertTrue((figure / "figure.pdf").is_file())
                render_archived_figure(figure, "dispatch_peak" if chart.startswith("02")
                                       else "dispatch_year")
                self.assertTrue((figure / "renderer.json").is_file())
            self.assertLess(run.summary["max_balance_residual_mw"], 1e-7)

    def test_leap_curve_rejects_misaligned_day_labels(self):
        """8760个值仍可能把闰日错放进2月；导入前需核对日期顺序。"""
        with TemporaryDirectory() as folder:
            year_folder = Path(folder) / "2028"
            curve_folder = year_folder / "CURVE"
            curve_folder.mkdir(parents=True)
            book = Workbook()
            sheet = book.active
            sheet.append(["日期"] + list(range(1, 25)))
            day = date(2028, 1, 1)
            while day.year == 2028:
                if day != date(2028, 2, 29):
                    sheet.append([day.strftime("%Y/%m/%d")] + [0.5] * 24)
                day += timedelta(days=1)
            wrong_row = (date(2028, 3, 1) - date(2028, 1, 1)).days + 1
            sheet.cell(wrong_row, 1).value = "2028/02/29"
            path = curve_folder / "synthetic.xlsx"
            book.save(path)
            with self.assertRaisesRegex(ValueError, "闰年曲线"):
                _read_curve(year_folder, "synthetic", {})
            sheet.cell(wrong_row, 1).value = "2028/03/01"
            book.save(path)
            book.close()
            self.assertEqual(len(_read_curve(year_folder, "synthetic", {})), 8760)

    def test_cost_duplicate_and_curve_overrating_are_rejected(self):
        """重复成本行不能静默覆盖；单位出力>1会突破额定MW，须拦截。"""
        with TemporaryDirectory() as folder:
            root = Path(folder)
            csv_path = root / "cost.csv"
            csv_path.write_text("technology,annual_capital_yuan_per_mw,variable_yuan_per_mwh\n" +
                                "".join(f"{kind},100,10\n" for kind in TECH_LABELS) +
                                "coal,999,999\n", encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "重复技术"):
                read_costs(csv_path)
            year_folder = root / "2029"
            curve_folder = year_folder / "CURVE"
            curve_folder.mkdir(parents=True)
            book = Workbook(write_only=True)
            sheet = book.create_sheet()
            sheet.append(["日期"] + list(range(1, 25)))
            for day in range(365):
                # 第1日第1小时1.2倍额定值，其余小时在合法单位值范围内。
                sheet.append([f"day-{day + 1}"] +
                             ([1.2] + [0.5] * 23 if day == 0 else [0.5] * 24))
            book.save(curve_folder / "overrated.xlsx")
            with self.assertRaisesRegex(ValueError, "0–1范围"):
                _read_curve(year_folder, "overrated", {})

    def test_storage_zero_charge_ratio_is_rejected(self):
        """储能充电倍率0不得被默认值1替代，否则会凭空增加充电能力。"""
        station = [None] * 16
        station[1], station[4], station[7], station[8] = "SyntheticStorage", 385, 0, 0.81
        unit = [None] * 9
        unit[8] = 2  # 库容Emax=P额定×2h；充电倍率仍必须显式有效。
        costs = {"battery": {"variable_yuan_per_mwh": "0",
                             "annual_capital_yuan_per_mw": "100"}}
        with self.assertRaisesRegex(ValueError, "充电功率倍率"):
            _make_plant(1, (station, unit, 10), 10, Path("."), {}, costs, {}, False)
        station[4] = 300  # 常规机组最小出力比例也必须落在[0,1]。
        unit[7], unit[8] = 1.5, 0
        costs["coal"] = costs["battery"]
        with self.assertRaisesRegex(ValueError, "最小出力比例"):
            _make_plant(2, (station, unit, 10), 10, Path("."), {}, costs, {}, False)

    def test_online_reserve_is_planning_feasibility_gate(self):
        """无缺供但仅余2MW备用、需3MW备用时，不能报告为可行规划。"""
        case = ProjectCase(2025, 2029,
                           (CasePlant(1, "Coal", "coal", 102, 0, 300, 0),),
                           (), (100.0,) * 8760, (0.0,) * 8760,
                           0.0, 0, (), "", "")
        score = evaluate_project(case, ())
        self.assertAlmostEqual(score.unserved_mwh, 0.0)
        self.assertGreater(score.reserve_shortfall_mwh, 0.0)
        self.assertFalse(score.feasible)

    def test_charging_storage_is_not_counted_as_upward_reserve(self):
        """充电10MW时不能再把刚充入的内部电量作为同小时备用。"""
        case = ProjectCase(2025, 2029,
                           (CasePlant(1, "Coal", "coal", 100, 1.0, 300, 0),
                            CasePlant(2, "Battery", "battery", 20, 0, 0, 0,
                                      (), 2, 1, 0.81)),
                           (), (90.0,) * 8760, (0.0,) * 8760,
                           0.0, 0, (), "", "")
        run = simulate_project(case, (), keep_hourly=True)
        self.assertAlmostEqual(run.hourly[0].storage_charge_mw, 10.0)
        self.assertAlmostEqual(run.hourly[0].reserve_available_mw, 0.0)
        self.assertAlmostEqual(run.hourly[0].reserve_shortfall_mw, 2.7)


if __name__ == "__main__":
    unittest.main()
