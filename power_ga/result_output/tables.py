"""结果输出：把所有运行表合入一个XLSX，同时保留CSV与Origin工程。"""

import csv
import json
from dataclasses import asdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..common.domain import PLANNING_MONTH_HOURS, Problem
from ..initial_population.encoding import decode
from ..genetic_iteration.solver import GAResult
from ..operation_simulation import simulate
from .origin_plot import plot_results


def _write_csv(path: Path, header, rows):
    """把可供脚本再次读取的逐表CSV保存为Windows Excel兼容编码。"""
    # utf-8-sig 方便在 Windows Excel 中直接打开中文表格。
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(header)
        writer.writerows(rows)


def _spreadsheet_value(value):
    """CSV读回时恢复数值类型；Excel中的MW/MWh须能排序和计算。"""
    if not isinstance(value, str) or not value:
        return value
    try:
        number = float(value)
        return int(number) if number.is_integer() and len(value) < 16 else number
    except ValueError:
        return value


def _add_sheet(book, name, header, rows, numeric_digits=3):
    """同一工作簿内建立可筛选的平面表，首行固定为中文字段名。"""
    sheet = book.create_sheet(name)
    sheet.sheet_view.showGridLines = False
    sheet.append(header)
    for row in rows:
        sheet.append([_spreadsheet_value(value) for value in row])
    sheet.freeze_panes = "B2" if len(header) > 5 else "A2"
    sheet.auto_filter.ref = sheet.dimensions
    # 标题行深蓝配色；数值仍保存为Excel数值，不因显示格式而四舍五入。
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.font = Font(name="等线", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 28
    for col, label in enumerate(header, 1):
        sheet.column_dimensions[get_column_letter(col)].width = min(34, max(12, len(label) * 2 + 4))
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, float):
                cell.number_format = "#,##0." + "0" * numeric_digits
    return sheet


def _period_rows(hourly, period_lengths, label):
    """逐小时功率按Δt=1h积分：E=ΣP_h×1h；跨日储能库存取首末点。"""
    fields = ("load_mw", "thermal_mw", "wind_mw", "solar_mw", "storage_charge_mw",
              "storage_discharge_mw", "wind_curtailment_mw", "solar_curtailment_mw",
              "unserved_mw", "reserve_shortfall_mw", "overgeneration_mw")
    header = [label, "起始小时", "结束小时", "最大负荷MW", "期初库存MWh", "期末库存MWh"] + [
        title + "MWh" for title in ("负荷电量", "火电", "风电", "光伏", "储能充电", "储能放电",
                               "弃风", "弃光", "缺供", "备用不足积分", "不可消纳过剩")]
    rows, first = [], 0
    for number, length in enumerate(period_lengths, 1):
        part = hourly[first:first + length]
        rows.append([number, first, first + length - 1, max(h.load_mw for h in part),
                     part[0].storage_energy_start_mwh, part[-1].storage_energy_end_mwh] +
                    [sum(getattr(h, field) for h in part) for field in fields])
        first += length
    if first != len(hourly):
        raise ValueError("日/月统计时段没有覆盖全部8760小时")
    return header, rows


def _write_workbook(output_dir: Path, problem: Problem, result: GAResult, simulation):
    """把各类运行结果合入同一个工作簿，逐时原值不降采样。"""
    if simulation.hourly is None or len(simulation.hourly) != 8760:
        raise ValueError("全年工作簿要求完整8760小时模拟结果")
    book = Workbook()
    book.remove(book.active)
    # 年度总览保持费用、缺电、风光、储能、备用等原结果值。
    with (output_dir / "summary.csv").open(encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.reader(stream))
    _add_sheet(book, "年度总览", rows[0], rows[1:], 9)
    with (output_dir / "projects.csv").open(encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.reader(stream))
    _add_sheet(book, "规划项目", rows[0], rows[1:])

    # 逐时主表与CSV一一对应；小时0至8759均保留，不使用典型日替代。
    hour_headers = ["小时", "负荷MW", "风电可用MW", "光伏可用MW", "风电实际MW", "光伏实际MW",
                    "火电MW", "弃风MW", "弃光MW", "缺供MW", "储能充电MW", "储能放电MW",
                    "期初库存MWh", "期末库存MWh", "在线火电容量MW", "备用需求MW", "可用备用MW",
                    "备用缺口MW", "不可消纳过剩MW", "功率平衡残差MW"]
    _add_sheet(book, "8760小时平衡", hour_headers,
               (list(asdict(h).values()) for h in simulation.hourly), 6)
    if simulation.storage_hourly:
        # 一站一行一小时；SOC=E_t/Emax，库存守恒残差用于核对充放效率。
        _add_sheet(book, "储能逐时", ["小时", "储能电站", "充电MW", "放电MW", "期初库存MWh",
                                  "期末库存MWh", "SOC期末", "状态残差MWh"],
                   (list(asdict(h).values()) for h in simulation.storage_hourly), 6)
    daily_header, daily = _period_rows(simulation.hourly, [24] * 365, "日序号")
    _add_sheet(book, "逐日统计", daily_header, daily)
    # 教学算例也使用与项目流程相同的365日规划年月份边界。
    monthly_header, monthly = _period_rows(simulation.hourly,
                                           PLANNING_MONTH_HOURS, "月序号")
    _add_sheet(book, "逐月统计", monthly_header, monthly)
    if result.history:
        with (output_dir / "history.csv").open(encoding="utf-8-sig", newline="") as stream:
            rows = list(csv.reader(stream))
        _add_sheet(book, "遗传迭代", rows[0], rows[1:], 6)
    if simulation.solver_log:
        keys = list(simulation.solver_log[0])
        _add_sheet(book, "优化窗口", keys,
                   ([record[key] for key in keys] for record in simulation.solver_log), 6)
    _add_sheet(book, "口径说明", ["项目", "计算口径或边界"], [
        ["演示数据", "电站和负荷曲线均为合成算例；小时0表示年初第1小时。"],
        ["时间积分", "各小时长度均为1h；年度电量E=Σ(逐时功率MW×1h)。"],
        ["逐时平衡", "火电+风电实际+光伏实际+储能放电+缺供=负荷+储能充电+不可消纳过剩。"],
        ["储能状态", "E(t+1)=E(t)+充电功率×充电效率−放电功率/放电效率，跨日连续。"],
        ["备用", "备用缺口=max(0,需求−可用备用)；逐日和逐月列是功率×1h的积分。"],
        ["模式", problem.dispatch_mode + "；规则模式为透明基线，优化模式为滚动MILP扩展。"],
    ])
    book.save(output_dir / "全年运行结果汇总.xlsx")


def write_results(problem: Problem, result: GAResult, output_dir: Path, make_figures=True, simulation=None):
    """表格集中保存到tables，图片和Origin工程集中保存到figures。"""
    output_dir.mkdir(parents=True, exist_ok=True)
    # CSV及唯一汇总工作簿属于表格交付；图形输出在figures，不与表格混放。
    tables_dir = output_dir / "tables"
    tables_dir.mkdir(parents=True, exist_ok=True)
    best = result.best
    simulation = simulation if simulation is not None else simulate(problem, best.chromosome, keep_hourly=True)
    _write_csv(tables_dir / "summary.csv", ["指标", "数值", "单位"], [
        ["是否满足简化约束", "是" if best.feasible else "否", ""],
        ["新增电站", "、".join(decode(problem, best.chromosome)) or "无", ""],
        ["年化新增固定成本", f"{best.annual_fixed_cost_yuan:.2f}", "元/年"],
        ["年运行变动成本", f"{best.variable_cost_yuan:.2f}", "元/年"],
        ["年总成本", f"{best.total_cost_yuan:.2f}", "元/年"],
        ["全年缺供电量", f"{best.unserved_mwh:.9f}", "MWh"],
        ["缺供小时数", best.shortage_hours, "小时"],
        ["可信容量", f"{best.firm_capacity_mw:.3f}", "MW"],
        ["所需可信容量", f"{best.required_firm_capacity_mw:.3f}", "MW"],
        ["风电供电量", f"{simulation.wind_mwh:.3f}", "MWh"],
        ["风电可用电量", f"{simulation.wind_available_mwh:.3f}", "MWh"],
        ["光伏供电量", f"{simulation.solar_mwh:.3f}", "MWh"],
        ["光伏可用电量", f"{simulation.solar_available_mwh:.3f}", "MWh"],
        ["火电供电量", f"{simulation.thermal_mwh:.3f}", "MWh"],
        ["弃风电量", f"{simulation.wind_curtailment_mwh:.3f}", "MWh"],
        ["弃光电量", f"{simulation.solar_curtailment_mwh:.3f}", "MWh"],
        ["储能年充电量", f"{simulation.storage_charge_mwh:.3f}", "MWh"],
        ["储能年放电量", f"{simulation.storage_discharge_mwh:.3f}", "MWh"],
        ["储能年初库存", f"{simulation.storage_initial_mwh:.3f}", "MWh"],
        ["储能年末库存", f"{simulation.storage_final_mwh:.3f}", "MWh"],
        ["储能转换损失", f"{simulation.storage_loss_mwh:.3f}", "MWh"],
        ["运行备用不足积分", f"{simulation.reserve_shortfall_mwh:.9f}", "MW·h"],
        ["运行备用不足小时数", simulation.reserve_shortfall_hours, "小时"],
        ["最小出力不可消纳过剩", f"{simulation.overgeneration_mwh:.9f}", "MWh"],
        ["最大储能状态残差", f"{simulation.max_storage_residual_mwh:.9f}", "MWh"],
        ["最大功率平衡残差", f"{simulation.max_balance_residual_mw:.9f}", "MW"],
        ["终止原因", result.stop_reason, ""],
        ["实际评价的不同方案数", result.evaluations, "个"],
    ])
    _write_csv(tables_dir / "projects.csv",
               ["电站", "技术", "是否建设", "装机MW", "可信容量MW", "年化固定成本元"],
               ([p.name, p.technology, bit, p.capacity_mw, p.firm_capacity_mw,
                 p.annual_fixed_cost_yuan] for bit, p in zip(best.chromosome, problem.candidates)))
    _write_csv(tables_dir / "history.csv",
               ["代数", "历代最优年总成本元", "最优方案可行", "本代可行个体数", "最优方案违约程度",
                "本代最优成本元", "可行个体平均成本元", "可行个体中位成本元", "不同染色体数", "历代最优染色体"],
               ([row.generation, row.best_cost_yuan, int(row.best_feasible),
                 row.feasible_count, row.best_violation, row.generation_best_cost_yuan,
                 row.mean_cost_yuan, row.median_cost_yuan, row.unique_count,
                 "".join(map(str, row.best_chromosome))] for row in result.history))
    # 列名直接与只读结果结构对应，其他模块可按字段名读取，不依赖列位置。
    _write_csv(tables_dir / "hourly_dispatch.csv", list(asdict(simulation.hourly[0])),
               (list(asdict(h).values()) for h in simulation.hourly))
    if simulation.storage_hourly:
        _write_csv(tables_dir / "storage_hourly.csv", list(asdict(simulation.storage_hourly[0])),
                   (list(asdict(h).values()) for h in simulation.storage_hourly))
    if simulation.solver_log:
        (tables_dir / "solver_windows.json").write_text(json.dumps(simulation.solver_log, indent=2), encoding="utf-8")
    _write_workbook(tables_dir, problem, result, simulation)
    if make_figures:
        plot_results(result, simulation, output_dir, problem)
