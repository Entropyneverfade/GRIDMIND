"""多年度动态规划结果归档；与最终年份8760小时工作簿分开。"""

import csv
import json
from pathlib import Path

from openpyxl import Workbook

from ..data_collection.project_case import TECH_LABELS
from ..initial_population.encoding import STAGE_SEED_POLICY


def write_dynamic_report(result, case, output: Path, load_factor: float,
                         stage_input_mode="scaled", ga_comparison=None,
                         exogenous_projects=()):
    """保存阶段状态、建站位与费用，便于EXE下次直接载入画规划图。"""
    tables = Path(output) / "tables"
    tables.mkdir(parents=True, exist_ok=True)
    headers = ("规划年", "累计状态整数", "当年新增位掩码", "当年新增电站",
               "年度费用元", "折现费用元", "最大功率平衡残差MW")
    rows = []
    for stage in result.stages:
        names = "+".join(plant.name for bit, plant in enumerate(case.candidates)
                         if stage.newly_built & (1 << bit)) or "无"
        rows.append((stage.year, stage.state, stage.newly_built, names,
                     stage.annual_cost_yuan, stage.discounted_cost_yuan,
                     stage.max_balance_residual_mw))
    with (tables / "dynamic_stages.csv").open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(headers)
        writer.writerows(rows)
    book = Workbook()
    sheet = book.active
    sheet.title = "动态规划阶段"
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    # 源库独有项目不属于DP可选位；单列容量与年费，避免被阶段汇总掩盖。
    exogenous_rows = [(stage_year, plant.station_id, plant.name,
                       TECH_LABELS[plant.technology], plant.capacity_mw,
                       plant.annual_cost_yuan)
                      for stage_year, plant in exogenous_projects]
    if exogenous_rows:
        headers_exo = ("规划年", "电站ID", "电站", "技术", "外生投运容量MW", "年化资本费元")
        exo_sheet = book.create_sheet("年度外生计划")
        exo_sheet.append(headers_exo)
        for row in exogenous_rows:
            exo_sheet.append(row)
        with (tables / "exogenous_projects.csv").open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.writer(stream)
            writer.writerow(headers_exo)
            writer.writerows(exogenous_rows)
    comparison_audit = None
    if ga_comparison is not None:
        ga_result, ga_path = ga_comparison
        ga_feasible = bool(ga_result.best.feasible)
        # 只有GA找到可行路径时，费用差才有优化意义；否则留空并显示不可行。
        absolute_gap = (ga_result.best.total_cost_yuan - result.objective_yuan
                        if ga_feasible else None)
        gap = ((absolute_gap / result.objective_yuan) if ga_feasible and
               result.objective_yuan > 0 else
               (0.0 if ga_feasible and abs(absolute_gap) <= 0.01 else None))
        if ga_feasible and ga_result.best.total_cost_yuan < result.objective_yuan - 0.01:
            raise AssertionError("同口径GA费用低于DP最优值，请复查阶段费用或约束")
        comparison_audit = {
            "ga_feasible": ga_feasible,
            "ga_seed_policy": STAGE_SEED_POLICY,
            "ga_states": list(ga_path),
            "ga_objective_yuan": ga_result.best.total_cost_yuan if ga_feasible else None,
            "dp_objective_yuan": result.objective_yuan,
            "absolute_gap_yuan": absolute_gap,
            "relative_gap": gap,
            "ga_evaluated_chromosomes": ga_result.evaluations,
            "ga_stop_reason": ga_result.stop_reason}
        comparisons = [("算法", "可行", "累计阶段状态", "折现总费用元", "相对DP差距"),
                       ("GA", ga_feasible, str(ga_path),
                        ga_result.best.total_cost_yuan if ga_feasible else None, gap),
                       ("DP+Gurobi", True, str(tuple(s.state for s in result.stages)),
                        result.objective_yuan, 0.0)]
        compare_sheet = book.create_sheet("GA与DP对比")
        for row in comparisons:
            compare_sheet.append(row)
        with (tables / "ga_dp_comparison.csv").open("w", encoding="utf-8-sig", newline="") as stream:
            csv.writer(stream).writerows(comparisons)
        if len(ga_path) != len(result.stages):
            raise ValueError("GA建设路径与动态规划阶段数不一致")
        # GA与DP查同一状态费用表；逐年折现费为c_t(s)/(1+r)^t。
        stage_headers = ("规划年", "GA累计状态", "DP累计状态", "GA阶段可行",
                         "GA年度费用元", "DP年度费用元", "GA折现费用元",
                         "DP折现费用元", "GA减DP折现差元")
        stage_rows = []
        for index, (stage, ga_state) in enumerate(zip(result.stages, ga_path)):
            item = result.evaluations[index][ga_state]
            ga_discounted = item.total_cost_yuan / (1 + result.discount_rate) ** index
            stage_rows.append((stage.year, ga_state, stage.state, item.feasible,
                               item.total_cost_yuan, stage.annual_cost_yuan,
                               ga_discounted, stage.discounted_cost_yuan,
                               ga_discounted - stage.discounted_cost_yuan))
        # 可行GA路径的逐年费用必须加总回GA报告值，否则说明编码或折现口径漂移。
        if ga_feasible and abs(sum(row[6] for row in stage_rows) -
                               ga_result.best.total_cost_yuan) > 0.01:
            raise AssertionError("GA逐年折现费与总费用不一致")
        stage_sheet = book.create_sheet("逐年GA与DP")
        stage_sheet.append(stage_headers)
        for row in stage_rows:
            stage_sheet.append(row)
        with (tables / "ga_dp_stage_comparison.csv").open(
                "w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.writer(stream)
            writer.writerow(stage_headers)
            writer.writerows(stage_rows)
        # 留存逐代可行个体与最优费用，方便判断GA何时达到或未达到DP下界。
        with (tables / "multiyear_ga_history.csv").open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.writer(stream)
            writer.writerow(("代", "历代最优折现费用元", "历代最优可行", "本代可行个体数", "染色体多样性"))
            writer.writerows((h.generation, h.best_cost_yuan, h.best_feasible,
                              h.feasible_count, h.unique_count) for h in ga_result.history)
    notes = book.create_sheet("模型口径")
    description = ("各阶段读取对应年度源库的存量、负荷、外送、风光和水文；仅当年出现的非终期增量按外生计划计费"
                   if stage_input_mode == "source" else
                   "前一年复用目标年风光曲线并缩放负荷；这是情景预测，不是年度源库重放")
    for item in (("项目", "说明"),
                 ("算法", "按年份Bellman状态递推；Gurobi 0/1状态选择模型独立核验"),
                 ("多年度GA初代", "全不建/全建后按规划年交织单项目种子；策略标识" + STAGE_SEED_POLICY),
                 ("阶段输入方式", stage_input_mode),
                 ("前一年负荷系数", load_factor if stage_input_mode == "scaled" else "不适用"),
                 ("曲线", description),
                 ("费用", "年化建设费用+8760小时运行变动费；次年折现"),
                 ("物理边界", "单区域确定性运行，不含网络潮流与概率可靠性")):
        notes.append(item)
    book.save(tables / "动态规划阶段汇总.xlsx")
    # 不重复写入源路径；源文件哈希由最终年运行存档清单保存。
    (Path(output) / "dynamic_audit.json").write_text(json.dumps({
        "algorithm": "dynamic_programming", "years": [s.year for s in result.stages],
        "states": [s.state for s in result.stages],
        "objective_yuan": result.objective_yuan,
        "gurobi_objective_yuan": result.gurobi_objective_yuan,
        "discount_rate": result.discount_rate,
        "stage_input_mode": stage_input_mode,
        "exogenous_projects_count": len(exogenous_rows),
        "exogenous_annual_cost_yuan": sum(row[-1] for row in exogenous_rows),
        "previous_load_factor": load_factor if stage_input_mode == "scaled" else None,
        "max_balance_residual_mw": max(s.max_balance_residual_mw for s in result.stages),
        "ga_comparison": comparison_audit
    }, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
