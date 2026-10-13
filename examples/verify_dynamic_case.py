"""独立回算公开合成GA/DP案例；不读取用户的项目库或私有结果。"""

import argparse
import csv
import json
from math import isclose
from pathlib import Path

from openpyxl import load_workbook

from examples.synthetic_dynamic_case import build_stages
from power_ga.common.domain import PLANNING_MONTH_HOURS
from power_ga.genetic_iteration.project_solver import evaluate_project
from power_ga.operation_simulation.project_engine import simulate_project


TECHNOLOGIES = ("coal", "gas", "hydro", "biomass", "pumped_discharge",
                "battery_discharge", "wind", "solar")


def _rows(path):
    """CSV读回原始数值文本，避免通过绘图或工作簿中的显示舍入验收。"""
    with path.open(encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream))


def _close(actual, expected, label, tolerance=0.01):
    """费用按分比较；物理残差另用更严格的MW/MWh绝对界。"""
    if not isclose(actual, expected, rel_tol=1e-10, abs_tol=tolerance):
        raise AssertionError(f"{label}不一致：{actual} != {expected}")


def verify(destination: Path):
    """用公开固定种子重跑所选路径，核对阶段账本、守恒及八类实际出力。"""
    destination = Path(destination)
    audit = json.loads((destination / "synthetic_dynamic_audit.json").read_text(encoding="utf-8"))
    if audit.get("synthetic") is not True or len(audit["stages"]) not in (2, 3):
        raise AssertionError("仅支持附带合成标志的两年度或三年度案例")
    cases = build_stages(full_horizon=len(audit["stages"]) == 3)
    if tuple(audit["stages"]) != tuple(case.target_year for case in cases):
        raise AssertionError("存档年份与固定种子合成输入不一致")
    stage_rows = _rows(destination / "tables" / "stage_comparison.csv")
    state_rows = _rows(destination / "tables" / "all_stage_states.csv")
    summary = {row["指标"]: float(row["数值"])
               for row in _rows(destination / "tables" / "summary.csv")}
    if len(stage_rows) != len(cases) or len(state_rows) != len(cases) * 2 ** len(cases[0].candidates):
        raise AssertionError("年度行数或完整状态费用表行数错误")
    rate = summary["折现率"]
    _close(summary["DP_Bellman_元"], audit["dp_cost_yuan"], "DP总费用")
    _close(summary["DP_Gurobi_元"], audit["gurobi_cost_yuan"], "Gurobi总费用")
    _close(summary["GA_元"], audit["ga_cost_yuan"], "GA总费用")
    _close(audit["dp_cost_yuan"], audit["gurobi_cost_yuan"], "Bellman/Gurobi总费用")
    state_lookup = {(int(row["规划年"]), int(row["状态整数"])): row for row in state_rows}
    physical, technologies = [], []
    dp_total = ga_total = 0.0
    previous = 0
    for stage, (case, row) in enumerate(zip(cases, stage_rows)):
        year = int(row["规划年"])
        state = int(row["DP累计状态整数"])
        ga_state = int(row["GA累计状态整数"])
        if year != case.target_year or state & previous != previous:
            raise AssertionError("规划年错位或累计项目发生退出")
        if int(row["DP本年新增位掩码"]) != state & ~previous:
            raise AssertionError("本年新增项目位掩码错误")
        for name, selected in (("DP", state), ("GA", ga_state)):
            item = state_lookup[(year, selected)]
            if int(item["可行"]) != 1:
                raise AssertionError(f"{year}年{name}方案不可行")
            _close(float(row[f"{name}阶段年费用元"]), float(item["年总费用元"]),
                   f"{year}年{name}状态费用")
            _close(float(row[f"{name}折现费用元"]),
                   float(item["年总费用元"]) / (1 + rate) ** stage,
                   f"{year}年{name}折现费用")
        # 不读取求解缓存：用固定种子输入独立重跑被选中的8760小时方案。
        bits = tuple((state >> bit) & 1 for bit in range(len(case.candidates)))
        evaluation = evaluate_project(case, bits)
        operation = simulate_project(case, bits, keep_hourly=True)
        _close(evaluation.total_cost_yuan, float(row["DP阶段年费用元"]),
               f"{year}年独立回算费用")
        if not evaluation.feasible or len(operation.hourly) != 8760:
            raise AssertionError(f"{year}年物理仿真不可行或小时数错误")
        balance = max(abs(hour.balance_residual_mw) for hour in operation.hourly)
        storage = max((abs(item[-1]) for item in operation.storage_hourly), default=0.0)
        if balance > 1e-8 or storage > 1e-8:
            raise AssertionError(f"{year}年功率或储能状态不守恒")
        if any(item[7] < -1e-9 or item[7] > 1 + 1e-9 for item in operation.storage_hourly):
            raise AssertionError(f"{year}年储能SOC越界")
        # 月水量上限按已投运水电容量累计，未用水量允许为正。
        hydros = [plant for plant in case.existing + tuple(p for p, bit in zip(case.candidates, bits) if bit)
                  if plant.technology == "hydro"]
        month_offset = 0
        hydro_gap = 0.0
        for month, hours in enumerate(PLANNING_MONTH_HOURS):
            actual = sum(hour.hydro_mw for hour in operation.hourly[month_offset:month_offset + hours])
            budget = sum(p.capacity_mw * hours * 0.42 * p.hydro_monthly_factor[month]
                         for p in hydros)
            hydro_gap = max(hydro_gap, actual - budget)
            month_offset += hours
        if hydro_gap > 1e-6:
            raise AssertionError(f"{year}年水电月预算超限")
        physical.append((year, state, 8760, balance, storage, hydro_gap,
                         operation.summary["unserved"],
                         operation.summary["reserve_shortfall"],
                         operation.summary["overgeneration"]))
        technologies.append((year,) + tuple(operation.summary[name] for name in TECHNOLOGIES))
        dp_total += float(row["DP折现费用元"])
        ga_total += float(row["GA折现费用元"])
        previous = state
    _close(dp_total, audit["dp_cost_yuan"], "DP各年折现费用之和")
    _close(ga_total, audit["ga_cost_yuan"], "GA各年折现费用之和")
    if any(value <= 0 for value in technologies[-1][1:]):
        raise AssertionError("最终年未体现八类电源的实际供电或储能放电")
    book_name = "全规划年算法对比汇总.xlsx" if len(cases) == 3 else "双年度算法对比汇总.xlsx"
    book = load_workbook(destination / "tables" / book_name, read_only=True, data_only=True)
    try:
        sheet = book["年度建站与费用"]
        if sheet.max_row != len(cases) + 1:
            raise AssertionError("工作簿年度建站表与CSV行数不一致")
        for index, row in enumerate(stage_rows, start=2):
            _close(float(sheet.cell(index, 4).value), float(row["DP阶段年费用元"]),
                   f"工作簿第{index}行费用")
    finally:
        book.close()
    # 表格与图片分别保存；审计表只写合成数据，不含外部路径。
    tables = destination / "tables"
    for name, header, rows in (
        ("stage_physical.csv", ("规划年", "DP状态", "小时数", "最大功率残差MW",
                                "最大储能状态残差MWh", "最大水电月预算超额MWh",
                                "确定性缺供MWh", "在线备用不足MWh", "不可消纳过剩MWh"), physical),
        ("technology_generation.csv", ("规划年",) + TECHNOLOGIES, technologies)):
        with (tables / name).open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.writer(stream)
            writer.writerow(header)
            writer.writerows(rows)
    result = {"synthetic": True, "status": "PASS", "stages": len(cases),
              "hours_per_stage": 8760, "states_per_stage": 2 ** len(cases[0].candidates),
              "dp_states": audit["dp_states"], "ga_states": audit["ga_states"],
              "dp_gurobi_gap_yuan": audit["dp_cost_yuan"] - audit["gurobi_cost_yuan"],
              "ga_dp_gap_yuan": audit["ga_cost_yuan"] - audit["dp_cost_yuan"],
              "max_balance_residual_mw": max(row[3] for row in physical),
              "max_storage_residual_mwh": max(row[4] for row in physical),
              "max_hydro_monthly_excess_mwh": max(row[5] for row in physical),
              "final_year_technologies_positive": 8}
    (destination / "synthetic_dynamic_verification.json").write_text(
        json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return result


def main():
    """命令行指定已生成的公开合成目录；失败抛出异常并返回非零退出码。"""
    parser = argparse.ArgumentParser(description="独立核对公开合成多年度GA/DP案例")
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    print(json.dumps(verify(args.output), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
