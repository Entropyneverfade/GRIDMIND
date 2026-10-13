"""公开合成两年度规划：同一阶段费用表比较遗传算法与动态规划。"""

import argparse
import csv
import json
from dataclasses import replace
from pathlib import Path

from openpyxl import Workbook

from examples.synthetic_eight_type import build_case
from power_ga.common.domain import GASettings
from power_ga.dynamic_planning import compare_ga_on_stages, solve_stage_planning
from power_ga.operation_simulation.project_engine import simulate_project
from power_ga.result_output.project_report import write_project_report
from power_ga.result_output.origin_plot import plot_dynamic_stage_costs


def build_stages():
    """同一八类资源和气象曲线，首年负荷缩至82%，次年恢复原曲线。"""
    case = build_case(with_candidates=True)
    # 这两个年份是公开教学情景，并非任何地区的年度预测或历史序列。
    first = replace(case, target_year=2026,
                    domestic_load_mw=tuple(0.82 * value for value in case.domestic_load_mw))
    second = replace(case, target_year=2027)
    return first, second


def _write_csv(path, header, rows):
    """表格与图片分目录；CSV保留未舍入的模型原值便于复算。"""
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(header)
        writer.writerows(rows)


def execute(destination: Path, discount_rate=0.07, origin_figures=False):
    """DP/Gurobi求阶段最优，再在同一状态费用表上执行PPT风格GA。"""
    cases = build_stages()
    destination = Path(destination)
    tables_dir = destination / "tables"
    tables_dir.mkdir(parents=True, exist_ok=True)
    dp = solve_stage_planning(cases, discount_rate)
    ga_settings = GASettings(64, 80, 12, 20, 2, 0.85, 0.045, 19)
    ga, ga_path = compare_ga_on_stages(cases, dp.evaluations, discount_rate, ga_settings)
    if not ga.best.feasible:
        raise AssertionError("合成两年度GA未找到可行方案")
    # 同一状态费用表是同口径比较的关键：气象、负荷、运行算法和成本完全一致。
    gap = (ga.best.total_cost_yuan - dp.objective_yuan) / dp.objective_yuan
    if gap < -1e-8:
        raise AssertionError("GA费用低于已由Gurobi校验的DP最优费用")
    dp_states = tuple(stage.state for stage in dp.stages)
    stage_rows = []
    for t, (case, stage) in enumerate(zip(cases, dp.stages)):
        ga_state = ga_path[t]
        stage_rows.append((case.target_year, stage.state, stage.newly_built,
                           stage.annual_cost_yuan, stage.discounted_cost_yuan,
                           stage.max_balance_residual_mw, ga_state,
                           dp.evaluations[t][ga_state].total_cost_yuan,
                           dp.evaluations[t][ga_state].total_cost_yuan /
                           (1 + discount_rate) ** t))
    stage_header = ("规划年", "DP累计状态整数", "DP本年新增位掩码", "DP阶段年费用元",
                    "DP折现费用元", "DP最大功率平衡残差MW", "GA累计状态整数",
                    "GA阶段年费用元", "GA折现费用元")
    _write_csv(tables_dir / "stage_comparison.csv", stage_header, stage_rows)
    if origin_figures:
        # 静态科研图必须由Origin原生绘制；每图独立保存OPJU、源数据与QA。
        plot_dynamic_stage_costs([row[0] for row in stage_rows],
                                 [row[3] for row in stage_rows],
                                 [row[4] for row in stage_rows],
                                 destination / "figures")
    cost_rows = ((case.target_year, state, int(item.feasible),
                  item.annual_fixed_cost_yuan, item.variable_cost_yuan,
                  item.total_cost_yuan, item.unserved_mwh,
                  item.reserve_shortfall_mwh, item.reserve_gap_mw)
                 for case, table in zip(cases, dp.evaluations) for state, item in table.items())
    _write_csv(tables_dir / "all_stage_states.csv",
               ("规划年", "状态整数", "可行", "年化资本费元", "变动费元", "年总费用元",
                "缺供MWh", "在线备用不足MW·h", "规划备用缺口MW"), cost_rows)
    history_rows = ((row.generation, row.best_cost_yuan,
                     row.generation_best_cost_yuan, row.feasible_count, row.unique_count)
                    for row in ga.history)
    _write_csv(tables_dir / "ga_history.csv",
               ("代数", "历代最优折现费用元", "本代最优折现费用元", "可行个体数",
                "不同染色体数"), history_rows)
    summary_rows = (("DP_Bellman_元", dp.objective_yuan),
                    ("DP_Gurobi_元", dp.gurobi_objective_yuan),
                    ("GA_元", ga.best.total_cost_yuan),
                    ("GA相对DP差距", gap),
                    ("GA评价次数", ga.evaluations),
                    ("DP状态运行次数", dp.evaluated_states),
                    ("折现率", discount_rate))
    _write_csv(tables_dir / "summary.csv", ("指标", "数值"), summary_rows)
    book = Workbook()
    book.active.title = "算法对比"
    book.active.append(("指标", "数值"))
    for row in summary_rows:
        book.active.append(row)
    for title, header, filename in (("年度建站与费用", stage_header, "stage_comparison.csv"),
                                    ("GA收敛", ("代数", "历代最优折现费用元", "本代最优折现费用元",
                                               "可行个体数", "不同染色体数"), "ga_history.csv"),
                                    ("全状态费用", ("规划年", "状态整数", "可行", "年化资本费元", "变动费元",
                                               "年总费用元", "缺供MWh", "在线备用不足MW·h",
                                               "规划备用缺口MW"), "all_stage_states.csv")):
        sheet = book.create_sheet(title)
        sheet.append(header)
        with (tables_dir / filename).open(encoding="utf-8-sig", newline="") as stream:
            for row in list(csv.reader(stream))[1:]:
                sheet.append([float(value) if value.replace(".", "", 1).replace("-", "", 1)
                              .replace("e", "", 1).replace("+", "", 1).isdigit() else value
                              for value in row])
    notes = book.create_sheet("模型口径")
    for row in (("项目", "说明"),
                ("来源", "固定种子公开合成八类电源；无外部项目数据"),
                ("阶段", "2026与2027两个连续规划年"),
                ("负荷", "首年=原合成8760曲线×0.82；次年=原曲线"),
                ("气象", "两年复用同一合成风光曲线，避免把随机天气差异混入算法比较"),
                ("状态", "8个候选项目累计0/1投运；已投运项目不能撤销"),
                ("费用", "各年=累计年化资本费+全年变动费；第2年按1/(1+r)折现"),
                ("求解", "Bellman递推，Gurobi 0/1状态模型独立校验；GA同表比较"),
                ("边界", "单区域确定性运行；无网络潮流和概率可靠性")):
        notes.append(row)
    book.save(tables_dir / "双年度算法对比汇总.xlsx")
    # 最终年DP方案仍以原报告接口保存完整8760小时/储能表，便于应用缩放查看。
    final_case = cases[-1]
    final_state = dp_states[-1]
    final_bits = tuple((final_state >> i) & 1 for i in range(len(final_case.candidates)))
    final_simulation = simulate_project(final_case, final_bits, keep_hourly=True)
    report_proxy = type("Result", (), {"best": dp.evaluations[-1][final_state], "history": ()})()
    write_project_report(final_case, report_proxy, final_simulation,
                         destination / "final_year", make_figures=False, synthetic=True)
    audit = {"synthetic": True, "stages": [case.target_year for case in cases],
             "candidate_bits": len(final_case.candidates), "dp_states": dp_states,
             "ga_states": ga_path, "dp_cost_yuan": dp.objective_yuan,
             "gurobi_cost_yuan": dp.gurobi_objective_yuan,
             "ga_cost_yuan": ga.best.total_cost_yuan, "ga_gap_fraction": gap,
             "max_balance_residual_mw": max(s.max_balance_residual_mw for s in dp.stages),
             "solver": "Gurobi 13.0.3 verified"}
    (destination / "synthetic_dynamic_audit.json").write_text(
        json.dumps(audit, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return audit


def main():
    """独立CLI明确输出目录；可用同一脚本逐项复现比较报表。"""
    parser = argparse.ArgumentParser(description="公开合成双年度动态规划与GA对比")
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--discount-rate", type=float, default=0.07)
    parser.add_argument("--origin-figures", action="store_true")
    args = parser.parse_args()
    print(json.dumps(execute(args.output, args.discount_rate, args.origin_figures),
                     ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
